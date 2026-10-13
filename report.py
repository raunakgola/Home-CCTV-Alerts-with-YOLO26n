"""The Excel report: one row per visit and a per-camera summary."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from helpers import fmt_dur
from zones import ZONES


def save_excel(rows, path):
    """Write the Visits and Summary sheets; returns (real path used, number of rows)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Visits"
    ws.append(["Camera", "Object", "Visit #", "First seen", "Last seen",
               "Duration (h:mm:ss)", "Duration (sec)", "Zone time (sec)", "Note"])
    summary = {}
    for v in rows:
        ws.append([v.cam, v.cls, v.id, v.first, v.last, fmt_dur(v.seconds()),
                   round(v.seconds(), 1), round(v.zone_seconds, 1) if v.cam in ZONES else "", v.note()])
        for col in (4, 5):
            ws.cell(row=ws.max_row, column=col).number_format = "yyyy-mm-dd hh:mm:ss"
        s = summary.setdefault((v.cam, v.cls), [0, 0.0, 0.0])
        s[0] += 1
        s[1] += v.seconds()
        s[2] += v.zone_seconds
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", (14, 13, 9, 21, 21, 18, 15, 16, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Camera", "Object", "Visits", "Total time (h:mm:ss)", "Total time (sec)",
                "Time in zone (sec)"])
    for (cam, cls), (n, secs, zsecs) in sorted(summary.items()):
        ws2.append([cam, cls, n, fmt_dur(secs), round(secs, 1), round(zsecs, 1) if cam in ZONES else ""])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for col, w in zip("ABCDEF", (14, 13, 9, 21, 17, 18)):
        ws2.column_dimensions[col].width = w
    try:
        wb.save(path)
    except PermissionError:      # the file is open in Excel
        path = path.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    return path, len(rows)
